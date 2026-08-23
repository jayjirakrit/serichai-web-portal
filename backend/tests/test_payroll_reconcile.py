import io

import numpy as np
import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from services.payroll_reconcile_service import (
    KNOWN_DEPARTMENTS,
    OFFICE_DEPARTMENT,
    PayrollReconcileRequestError,
    _build_employee_frame,
    _half_up_round2,
    _locate_header_positions,
    _read_raw_sheet,
    _recalculate_wages,
    _resolve_sheet_name,
    build_discrepancy_report,
    build_reconciliation_report,
    reconcile_payroll,
)

METAL_DEPARTMENT = "โรงเหล็ก"
SEWING_DEPARTMENT = "โรงเย็บ"
CONTRACT_DEPARTMENT = "พนง.ชั่วคราว"

PERIOD = "2026-08"
SHEET_NAME = "ประจำ 8-69"  # 'YYYY-MM' -> 'ประจำ {M-yy}' with a Buddhist Era (BE = Gregorian + 543) year

# Fixed column layout used by the fixture builder below, matching
# payroll_reconcile_service.py's FIXED_COLUMNS/OFFSET (research.md #2).
WORKDAY_COL = 5
WEEKEND_COL = 6
BASE_COL = 7
TOTAL_WORKDAY_HOUR_COL = BASE_COL + 1
LUMP_SUM_COL = BASE_COL + 6
DILIGENT_COL = BASE_COL + 10
SHIFT_FEE_COL = BASE_COL + 11
BONUS_COL = BASE_COL + 13
SOCIAL_SECURITY_COL = BASE_COL + 15
ADVANCE_COL = BASE_COL + 16
DEDUCT_LOAN_COL = BASE_COL + 17
TAX_COL = BASE_COL + 18
ACTUAL_TOTAL_WAGE_COL = BASE_COL + 19
ROW_WIDTH = ACTUAL_TOTAL_WAGE_COL + 1


def _blank_row() -> list:
    return [None] * ROW_WIDTH


def _attendance_row(department: str, emp: dict) -> list:
    row = _blank_row()
    row[0] = emp.get("index", "1")
    row[1] = department
    row[2] = emp.get("first_name", "")
    row[3] = emp.get("last_name", "")
    row[4] = emp.get("salary", 0)
    row[WORKDAY_COL] = emp.get("workday", 0)
    row[WEEKEND_COL] = emp.get("weekend", 0)
    row[TOTAL_WORKDAY_HOUR_COL] = emp.get("total_workday_hour", 0)
    row[LUMP_SUM_COL] = emp.get("lump_sum", 0)
    row[DILIGENT_COL] = emp.get("diligent_allowance", 0)
    row[SHIFT_FEE_COL] = emp.get("shift_fee", 0)
    row[BONUS_COL] = emp.get("bonus", 0)
    row[SOCIAL_SECURITY_COL] = emp.get("social_security", 0)
    row[ADVANCE_COL] = emp.get("advance_payment", 0)
    row[DEDUCT_LOAN_COL] = emp.get("deduct_loan_fund", 0)
    row[TAX_COL] = emp.get("tax", 0)
    row[ACTUAL_TOTAL_WAGE_COL] = emp.get("actual_total_wage", 0)
    return row


def _overtime_row(emp: dict) -> list:
    row = _blank_row()
    row[WORKDAY_COL] = emp.get("workday_ot", 0)
    row[WEEKEND_COL] = emp.get("weekend_ot", 0)
    return row


def _build_workbook(sections: list[dict], sheet_name: str = SHEET_NAME) -> bytes:
    """Builds a minimal `{sheet_name}` sheet matching the legacy batch's layout
    (data-model.md / research.md #2): a header row with the 'วันหยุด'/'วันปกติ'
    marker cells, followed by department-header rows each followed by paired
    attendance/overtime employee rows. `sections` is a list of
    `{"department": str, "employees": [emp_dict, ...]}`; an employee dict may
    set `no_ot: True` to omit its overtime row (data-error scenario)."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(["Company Header"])
    ws.append([])
    header_row = _blank_row()
    header_row[WEEKEND_COL] = "วันหยุด"
    header_row[BASE_COL] = "วันปกติ"
    ws.append(header_row)

    for section in sections:
        dept_row = _blank_row()
        dept_row[1] = section["department"]
        ws.append(dept_row)
        for emp in section["employees"]:
            ws.append(_attendance_row(section["department"], emp))
            if not emp.get("no_ot"):
                ws.append(_overtime_row(emp))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _find(discrepancies: list[dict], first_name: str) -> dict:
    return next(d for d in discrepancies if d["employee_name"] and first_name in d["employee_name"])


# --- Sheet-name resolution (research.md #3) -----------------------------


def test_resolve_sheet_name_uses_buddhist_era_m_yy_token():
    assert _resolve_sheet_name("2026-08") == "ประจำ 8-69"
    assert _resolve_sheet_name("2026-01") == "ประจำ 1-69"


# --- Header positions & department/row-pairing detection (research.md #2) ---


def test_header_positions_resolved_relative_to_total_workday_marker():
    content = _build_workbook([{"department": OFFICE_DEPARTMENT, "employees": [{"salary": 300, "actual_total_wage": 300}]}])
    raw = _read_raw_sheet(content, SHEET_NAME)
    header = _locate_header_positions(raw)

    assert header["weekend"] == WEEKEND_COL
    assert header["base"] == BASE_COL
    assert header["total_workday_hour"] == TOTAL_WORKDAY_HOUR_COL
    assert header["actual_total_wage"] == ACTUAL_TOTAL_WAGE_COL


def test_department_labels_ffill_and_attendance_overtime_rows_pair_up():
    content = _build_workbook(
        [
            {
                "department": OFFICE_DEPARTMENT,
                "employees": [
                    {"index": "1", "first_name": "หนึ่ง", "salary": 300, "actual_total_wage": 300},
                    {"index": "2", "first_name": "สอง", "salary": 300, "actual_total_wage": 300},
                ],
            }
        ]
    )
    raw = _read_raw_sheet(content, SHEET_NAME)
    header = _locate_header_positions(raw)
    employees = _build_employee_frame(raw, header)

    assert len(employees) == 2
    assert set(employees["department"]) == {OFFICE_DEPARTMENT}
    assert set(employees["employee_id"]) == {"1", "2"}
    assert employees["has_overtime_row"].all()


# --- Half-up rounding (research.md #4a) ---------------------------------


def test_half_up_round2_rounds_ties_away_from_zero_unlike_numpy_default():
    value = pd.Series([0.125])
    assert _half_up_round2(value).iloc[0] == pytest.approx(0.13)
    # numpy/pandas default rounding is round-half-to-even, and would give 0.12
    # here — the exact divergence this helper exists to avoid (research.md #4a).
    assert round(float(np.round(value.iloc[0], 2)), 2) == pytest.approx(0.12)


# --- Wage formulas & discrepancy detection (research.md #4) -------------


def test_all_matched_office_and_factory_employees_yield_zero_discrepancies():
    content = _build_workbook(
        [
            {"department": OFFICE_DEPARTMENT, "employees": [{"first_name": "ออฟฟิศ", "salary": 300, "actual_total_wage": 300}]},
            {
                "department": METAL_DEPARTMENT,
                "employees": [
                    {
                        "first_name": "โรงงาน",
                        "salary": 300,
                        "workday": 2,
                        "total_workday_hour": 16,
                        "actual_total_wage": 1200,
                    }
                ],
            },
        ]
    )

    result = reconcile_payroll(content, "payroll.xlsx", PERIOD)

    assert result["summary"]["total_employees_reconciled"] == 2
    assert result["summary"]["matched_count"] == 2
    assert result["summary"]["discrepancy_count"] == 0
    assert result["discrepancies"] == []
    assert result["discrepancy_report"] is None


def test_office_and_factory_discrepancies_use_their_own_department_formula():
    content = _build_workbook(
        [
            {"department": OFFICE_DEPARTMENT, "employees": [{"first_name": "ออฟฟิศ", "salary": 300, "actual_total_wage": 250}]},
            {
                "department": METAL_DEPARTMENT,
                "employees": [
                    {
                        "first_name": "โรงงาน",
                        "salary": 300,
                        "workday": 2,
                        "total_workday_hour": 16,
                        "actual_total_wage": 1300,
                    }
                ],
            },
        ]
    )

    result = reconcile_payroll(content, "payroll.xlsx", PERIOD)

    assert result["summary"]["discrepancy_count"] == 2
    office_entry = _find(result["discrepancies"], "ออฟฟิศ")
    assert office_entry["status"] == "discrepancy"
    assert office_entry["recalculated_amount"] == pytest.approx(300)
    assert office_entry["variance"] == pytest.approx(-50)

    factory_entry = _find(result["discrepancies"], "โรงงาน")
    assert factory_entry["status"] == "discrepancy"
    assert factory_entry["recalculated_amount"] == pytest.approx(1200)
    assert factory_entry["variance"] == pytest.approx(100)


def test_blank_total_workday_hour_cell_still_recalculates_factory_wage():
    """A blank 'ชม.ปกติ' cell (no partial-hour attendance that day) must be
    treated as 0, not left as NaN — NaN silently poisons the whole /8-based
    factory formula and makes a real discrepancy fall through to 'matched'
    (recalculated_amount stays None) instead of being flagged."""
    content = _build_workbook(
        [
            {
                "department": METAL_DEPARTMENT,
                "employees": [
                    {
                        "first_name": "โรงงาน",
                        "salary": 300,
                        "workday": 2,
                        "total_workday_hour": None,
                        "actual_total_wage": 1300,
                    }
                ],
            },
        ]
    )

    result = reconcile_payroll(content, "payroll.xlsx", PERIOD)

    entry = _find(result["discrepancies"], "โรงงาน")
    assert entry["status"] == "discrepancy"
    assert entry["recalculated_amount"] == pytest.approx(600)
    assert entry["variance"] == pytest.approx(700)


def test_unrecognized_department_flagged_without_formula_applied():
    content = _build_workbook(
        [{"department": "ทดสอบ", "employees": [{"first_name": "ไม่รู้จัก", "salary": 300, "actual_total_wage": 999}]}]
    )

    result = reconcile_payroll(content, "payroll.xlsx", PERIOD)

    entry = _find(result["discrepancies"], "ไม่รู้จัก")
    assert entry["status"] == "unrecognizedDepartment"
    assert entry["recalculated_amount"] is None
    assert entry["stated_amount"] == pytest.approx(999)


def test_missing_overtime_row_flagged_as_data_error():
    content = _build_workbook(
        [
            {
                "department": OFFICE_DEPARTMENT,
                "employees": [{"first_name": "ไม่มีโอที", "salary": 300, "actual_total_wage": 300, "no_ot": True}],
            }
        ]
    )

    result = reconcile_payroll(content, "payroll.xlsx", PERIOD)

    entry = _find(result["discrepancies"], "ไม่มีโอที")
    assert entry["status"] == "dataError"
    assert entry["recalculated_amount"] is None


# --- Request-level failures (FR-011) -------------------------------------


def test_period_with_no_matching_sheet_raises_request_error():
    content = _build_workbook([{"department": OFFICE_DEPARTMENT, "employees": [{"salary": 300, "actual_total_wage": 300}]}])

    with pytest.raises(PayrollReconcileRequestError):
        reconcile_payroll(content, "payroll.xlsx", "2099-01")


def test_missing_header_markers_raises_request_error():
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(["no header markers here"])
    buf = io.BytesIO()
    wb.save(buf)

    with pytest.raises(PayrollReconcileRequestError):
        reconcile_payroll(buf.getvalue(), "payroll.xlsx", PERIOD)


# --- Report generation (research.md #6, User Story 2) --------------------


def test_reconciliation_report_lists_every_employee_matched_and_flagged():
    content = _build_workbook(
        [
            {"department": OFFICE_DEPARTMENT, "employees": [{"first_name": "ตรงกัน", "salary": 300, "actual_total_wage": 300}]},
            {"department": OFFICE_DEPARTMENT, "employees": [{"index": "2", "first_name": "ไม่ตรง", "salary": 300, "actual_total_wage": 250}]},
        ]
    )
    result = reconcile_payroll(content, "payroll.xlsx", PERIOD)

    import base64

    report_bytes = base64.b64decode(result["reconciliation_report"]["content_base64"])
    ws = load_workbook(io.BytesIO(report_bytes)).active
    names = [ws.cell(r, 3).value for r in range(2, ws.max_row + 1)]
    assert "ตรงกัน" in names
    assert "ไม่ตรง" in names


def test_discrepancy_report_is_none_when_all_matched_and_populated_otherwise():
    all_matched = _build_workbook(
        [{"department": OFFICE_DEPARTMENT, "employees": [{"first_name": "ตรงกัน", "salary": 300, "actual_total_wage": 300}]}]
    )
    assert reconcile_payroll(all_matched, "payroll.xlsx", PERIOD)["discrepancy_report"] is None

    with_discrepancy = _build_workbook(
        [{"department": OFFICE_DEPARTMENT, "employees": [{"first_name": "ไม่ตรง", "salary": 300, "actual_total_wage": 250}]}]
    )
    result = reconcile_payroll(with_discrepancy, "payroll.xlsx", PERIOD)
    assert result["discrepancy_report"] is not None

    import base64

    report_bytes = base64.b64decode(result["discrepancy_report"]["content_base64"])
    ws = load_workbook(io.BytesIO(report_bytes)).active
    assert ws.cell(2, 3).value == "ไม่ตรง"
    assert ws.max_row == 2  # header + exactly the one discrepant employee


# --- Multi-department correctness (User Story 3) -------------------------


def test_discrepancies_across_all_four_departments_use_correct_formula_and_label():
    content = _build_workbook(
        [
            {"department": OFFICE_DEPARTMENT, "employees": [{"first_name": "ออฟฟิศ", "salary": 300, "actual_total_wage": 250}]},
            {
                "department": METAL_DEPARTMENT,
                "employees": [{"first_name": "เหล็ก", "salary": 300, "workday": 2, "total_workday_hour": 16, "actual_total_wage": 1300}],
            },
            {
                "department": SEWING_DEPARTMENT,
                "employees": [{"first_name": "เย็บ", "salary": 300, "workday": 2, "total_workday_hour": 16, "actual_total_wage": 1300}],
            },
            {
                "department": CONTRACT_DEPARTMENT,
                "employees": [{"first_name": "ชั่วคราว", "salary": 300, "workday": 2, "total_workday_hour": 16, "actual_total_wage": 1300}],
            },
        ]
    )

    result = reconcile_payroll(content, "payroll.xlsx", PERIOD)

    assert result["summary"]["discrepancy_count"] == 4
    for name, department in [
        ("ออฟฟิศ", OFFICE_DEPARTMENT),
        ("เหล็ก", METAL_DEPARTMENT),
        ("เย็บ", SEWING_DEPARTMENT),
        ("ชั่วคราว", CONTRACT_DEPARTMENT),
    ]:
        entry = _find(result["discrepancies"], name)
        assert entry["department"] == department

    office_entry = _find(result["discrepancies"], "ออฟฟิศ")
    assert office_entry["recalculated_amount"] == pytest.approx(300)  # office /30-based formula
    factory_style_amounts = [
        _find(result["discrepancies"], n)["recalculated_amount"] for n in ("เหล็ก", "เย็บ", "ชั่วคราว")
    ]
    assert factory_style_amounts == pytest.approx([1200, 1200, 1200])  # shared /8-based formula, all three categories
