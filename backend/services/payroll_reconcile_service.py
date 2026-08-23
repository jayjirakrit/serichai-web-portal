import base64
import io

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill

WEEKEND_MARKER = "วันหยุด"
TOTAL_WORKDAY_MARKER = "วันปกติ"
HEADER_ROW_INDEX = 2  # 0-based, matches the legacy batch's ReportHeaderPosition row

OFFICE_DEPARTMENT = "สำนักงาน"
METAL_DEPARTMENT = "โรงเหล็ก"
SEWING_DEPARTMENT = "โรงเย็บ"
CONTRACT_DEPARTMENT = "พนง.ชั่วคราว"
KNOWN_DEPARTMENTS = [OFFICE_DEPARTMENT, METAL_DEPARTMENT, SEWING_DEPARTMENT, CONTRACT_DEPARTMENT]

FIXED_COLUMNS = {
    "index": 0,
    "department": 1,
    "first_name": 2,
    "last_name": 3,
    "salary": 4,
    "workday_start": 5,
}

# Offsets relative to the "วันปกติ" (total-workday) column, mirroring the legacy
# batch's ReportHeaderPosition layout (research.md #2).
OFFSET = {
    "total_workday_hour": 1,
    "lump_sum": 6,
    "diligent_allowance": 10,
    "shift_fee": 11,
    "bonus": 13,
    "social_security": 15,
    "advance_payment": 16,
    "deduct_loan_fund": 17,
    "tax": 18,
    "actual_total_wage": 19,
}

REPORT_COLUMNS = [
    "#",
    "แผนก",
    "ชื่อ",
    "นามสกุล",
    "ค่าแรง",
    "วันปกติทั้งหมด",
    "ชม.ปกติทั้งหมด",
    "วันอาทิตย์ทั้งหมด",
    "วัน OT ปกติทั้งหมด",
    "วันOT วันหยุดทั้งหมด",
    "ค่าแรงปกติ",
    "ค่าจ้างเหงา",
    "รวม",
    "ค่าแรงวันอาทิตย์",
    "ค่าแรง OT",
    "เบี้ยขยัน",
    "ประกันสังคม",
    "รวมรับ",
    "เบิกเงินล่วงหน้า",
    "กยศ/บังคับคดี",
    "ภาษี",
    "สุทธิ",
    "คำนวนสุทธิ",
    "สถานะ",
    "รายละเอียด",
]

# Thai display labels for the reconciliation report only — the underlying
# `status`/`detail` values on each row stay in English (they're also
# returned as-is in the JSON API response, which the frontend types against).
STATUS_LABELS_TH = {
    "matched": "ตรงกัน",
    "discrepancy": "ไม่ตรงกัน",
    "unrecognizedDepartment": "ไม่พบแผนกในระบบ",
    "dataError": "ข้อมูลไม่ครบถ้วน",
}

DETAIL_LABELS_TH = {
    "dataError": "ไม่พบแถวข้อมูลโอที หรือข้อมูลเข้างานไม่ครบถ้วน ไม่สามารถคำนวณค่าแรงใหม่ได้",
    "unrecognizedDepartment": "แผนกนี้ไม่อยู่ในหมวดหมู่สูตรคำนวณค่าแรงที่รองรับ",
    "discrepancy": "ยอดค่าแรงสุทธิที่คำนวณใหม่ไม่ตรงกับยอดที่ระบุในไฟล์",
    "matched": "",
}


class PayrollReconcileRequestError(Exception):
    """Raised when the uploaded file, selected period, or file layout cannot be processed at all."""


# --- Sheet resolution & reading -----------------------------------------


def _resolve_sheet_name(period: str) -> str:
    """Converts 'YYYY-MM' to the legacy batch's sheet-name token, `ประจำ {M-yy}`
    with a Gregorian month (no leading zero) and a two-digit Buddhist Era year
    (BE = Gregorian year + 543), e.g. '2026-08' -> 'ประจำ 8-69'. The legacy batch
    is fed a BE-numbered year in its batchDate argument, so its `yy` token is
    always a BE year, not a Gregorian one — confirmed against real finance
    workbooks (e.g. 'ช69-1.xlsx' contains sheet 'ประจำ 1-69' for January 2026)."""
    try:
        year_str, month_str = period.split("-")
        year, month = int(year_str), int(month_str)
    except (ValueError, AttributeError) as exc:
        raise PayrollReconcileRequestError(f"Invalid period '{period}' — expected YYYY-MM.") from exc
    be_year = year + 543
    return f"ประจำ {month}-{be_year % 100:02d}"


def _read_raw_sheet(content: bytes, sheet_name: str) -> pd.DataFrame:
    try:
        raw = pd.read_excel(io.BytesIO(content), sheet_name=sheet_name, header=None, engine="openpyxl")
    except ValueError as exc:
        raise PayrollReconcileRequestError(f"Period '{sheet_name}' was not found in the uploaded file.") from exc
    except Exception as exc:
        raise PayrollReconcileRequestError("Payroll file is empty or unreadable.") from exc
    if raw.empty:
        raise PayrollReconcileRequestError("Payroll file is empty or unreadable.")
    return raw


# --- Header & department/row-pairing detection (research.md #2) --------


def _locate_header_positions(raw_df: pd.DataFrame) -> dict[str, int]:
    if HEADER_ROW_INDEX >= len(raw_df):
        raise PayrollReconcileRequestError("Payroll file is missing the expected header row.")

    header_row = raw_df.iloc[HEADER_ROW_INDEX]
    cells = [str(c).strip() if pd.notna(c) else "" for c in header_row.tolist()]

    weekend_col = next(
        (i for i in range(FIXED_COLUMNS["workday_start"], len(cells)) if cells[i] == WEEKEND_MARKER), None
    )
    if weekend_col is None:
        raise PayrollReconcileRequestError(f"Payroll file header is missing the '{WEEKEND_MARKER}' marker column.")

    base_col = next((i for i in range(weekend_col, len(cells)) if cells[i] == TOTAL_WORKDAY_MARKER), None)
    if base_col is None:
        raise PayrollReconcileRequestError(f"Payroll file header is missing the '{TOTAL_WORKDAY_MARKER}' marker column.")

    positions = dict(FIXED_COLUMNS)
    positions["weekend"] = weekend_col
    positions["base"] = base_col
    for name, offset in OFFSET.items():
        positions[name] = base_col + offset
    return positions


def _clean_str(value) -> str | None:
    if pd.isna(value):
        return None
    text = str(value).strip()
    return text or None


def _build_employee_frame(raw_df: pd.DataFrame, header: dict[str, int]) -> pd.DataFrame:
    """Splits the sheet's data rows into department-tagged attendance/overtime
    row pairs and merges them into one row per employee, per research.md #1-#2.
    A department-header row is detected structurally (its department name is
    written into the same column an attendance row uses for the employee's
    title, but — unlike an attendance row — its first-name cell is blank) so
    an unrecognized department name is still captured as a section, not
    silently skipped (research.md #5, FR-013). Real workbooks give a
    department-header row its own non-blank index (a per-section counter), so
    the index cell can't be used to detect it — see FIXED_COLUMNS note above."""
    data = raw_df.iloc[HEADER_ROW_INDEX + 1 :].reset_index(drop=True).copy()
    dept_col = header["department"]
    index_col = header["index"]
    first_name_col = header["first_name"]

    department = data[dept_col].apply(_clean_str)
    has_index = data[index_col].apply(_clean_str).notna()
    first_name_blank = data[first_name_col].apply(_clean_str).isna()
    # A department-header row states a section name but has no employee name
    # (structural, department-name-agnostic — an unrecognized name still opens
    # a section, research.md #5/FR-013). Every row below it, attendance or
    # overtime, is tagged with that department regardless of its own index cell.
    is_dept_header = department.notna() & first_name_blank

    data["department"] = department.where(is_dept_header).ffill()
    data["has_index"] = has_index
    data = data[~is_dept_header & data["department"].notna()].reset_index(drop=True)
    if data.empty:
        return data.assign(row_type=[], pair_id=[])

    data["row_seq"] = data.groupby("department").cumcount()

    # A department section ends at the first attendance-position row (even
    # row_seq) whose index cell is blank — mirrors the legacy batch's own
    # per-pair termination check (`StringUtils.hasText(index)`) and protects a
    # genuine overtime row (always odd row_seq, and often entirely blank for
    # an employee with zero overtime) from being mistaken for a trailing
    # filler row and dropped.
    bad_attendance = (data["row_seq"] % 2 == 0) & ~data["has_index"]
    if bad_attendance.any():
        cutoff = data.loc[bad_attendance].groupby("department")["row_seq"].min()
        data["cutoff"] = data["department"].map(cutoff)
        data = data[data["cutoff"].isna() | (data["row_seq"] < data["cutoff"])].reset_index(drop=True)
        data = data.drop(columns=["cutoff"])

    data["row_type"] = np.where(data["row_seq"] % 2 == 0, "attendance", "overtime")
    data["pair_id"] = data["row_seq"] // 2

    attendance = data[data["row_type"] == "attendance"].drop(columns=["row_seq", "row_type", "has_index"])
    overtime = data[data["row_type"] == "overtime"].drop(columns=["row_seq", "row_type", "has_index"])

    merged = attendance.merge(
        overtime,
        on=["department", "pair_id"],
        how="left",
        suffixes=("", "_ot"),
        indicator=True,
    )
    # pandas casts every column label to `str` once any column needs an "_ot"
    # suffix applied (even non-overlapping / non-suffixed ones), so every
    # positional column lookup below must go through this same string form.
    col = str

    merged["employee_id"] = merged[col(index_col)].apply(_clean_str)
    merged["first_name"] = merged[col(header["first_name"])].apply(_clean_str)
    merged["last_name"] = merged[col(header["last_name"])].apply(_clean_str)
    merged["salary"] = pd.to_numeric(merged[col(header["salary"])], errors="coerce")

    workday_cols = [col(c) for c in range(header["workday_start"], header["weekend"])]
    weekend_cols = [col(c) for c in range(header["weekend"], header["base"])]
    ot_workday_cols = [f"{c}_ot" for c in workday_cols]
    ot_weekend_cols = [f"{c}_ot" for c in weekend_cols]

    merged["has_overtime_row"] = merged["_merge"] == "both"
    merged["total_workday"] = merged[workday_cols].apply(pd.to_numeric, errors="coerce").sum(axis=1)
    merged["total_sunday"] = merged[weekend_cols].apply(pd.to_numeric, errors="coerce").sum(axis=1)
    merged["total_workday_ot_hour"] = np.where(
        merged["has_overtime_row"],
        merged[ot_workday_cols].apply(pd.to_numeric, errors="coerce").sum(axis=1),
        np.nan,
    )
    merged["total_weekend_ot_hour"] = np.where(
        merged["has_overtime_row"],
        merged[ot_weekend_cols].apply(pd.to_numeric, errors="coerce").sum(axis=1),
        np.nan,
    )

    return merged


# --- Wage formulas & discrepancy detection (research.md #4, #4a, #5) ---


def _half_up_round2(series: pd.Series) -> pd.Series:
    """Excel/Java RoundingMode.HALF_UP-equivalent rounding to 2 decimals for
    non-negative currency amounts (research.md #4a)."""
    return np.floor(series.astype(float) * 100 + 0.5) / 100


def _social_security_bracket(income: pd.Series) -> pd.Series:
    mid = np.floor(income * 0.05 + 0.5)
    return pd.Series(
        np.select(
            [income <= 0, income <= 1650, income <= 15000],
            [0.0, 83.0, mid],
            default=750.0,
        ),
        index=income.index,
    )


def _numeric_column(df: pd.DataFrame, col: int) -> pd.Series:
    # `_build_employee_frame`'s merge casts every column label to `str`
    # (see its `col = str` note) — mirror that here for positional lookups.
    return pd.to_numeric(df[str(col)], errors="coerce")


def _recalculate_wages(employees: pd.DataFrame, header: dict[str, int]) -> pd.DataFrame:
    df = employees.copy()

    is_office = df["department"] == OFFICE_DEPARTMENT
    is_known = df["department"].isin(KNOWN_DEPARTMENTS)

    salary = df["salary"]
    total_workday = df["total_workday"]
    total_sunday = df["total_sunday"]
    total_workday_ot = df["total_workday_ot_hour"]
    total_weekend_ot = df["total_weekend_ot_hour"]
    total_workday_hour = _numeric_column(df, header["total_workday_hour"]).fillna(0)
    lump_sum = _numeric_column(df, header["lump_sum"]).fillna(0)
    diligent_allowance = _numeric_column(df, header["diligent_allowance"]).fillna(0)
    shift_fee = _numeric_column(df, header["shift_fee"]).fillna(0)
    bonus = _numeric_column(df, header["bonus"]).fillna(0)
    stated_social_security = _numeric_column(df, header["social_security"]).fillna(0)
    advance_payment = _numeric_column(df, header["advance_payment"]).fillna(0)
    deduct_loan_fund = _numeric_column(df, header["deduct_loan_fund"]).fillna(0)
    tax = _numeric_column(df, header["tax"]).fillna(0)
    stated_actual_total_wage = _numeric_column(df, header["actual_total_wage"])

    # Office (สำนักงาน): salary-based with a /30 daily divisor.
    office_workday_salary = salary
    office_sunday_salary = _half_up_round2(salary * total_sunday / 30) * 2
    office_ot_salary = _half_up_round2(salary * total_workday_ot / 30 / 8 * 1.5) + _half_up_round2(
        salary * total_weekend_ot / 30 / 8 * 2
    )

    # Metal shop / sewing shop / contract staff: per-day-count based with a /8 hourly divisor.
    other_workday_salary = _half_up_round2(salary * total_workday) + _half_up_round2(salary * total_workday_hour / 8)
    other_sunday_salary = _half_up_round2(salary * total_sunday) * 2
    other_ot_salary = _half_up_round2(salary * total_workday_ot / 8 * 1.5) + _half_up_round2(
        salary * total_weekend_ot / 8 * 2
    )

    total_workday_salary = pd.Series(np.where(is_office, office_workday_salary, other_workday_salary), index=df.index)
    total_sunday_salary = pd.Series(np.where(is_office, office_sunday_salary, other_sunday_salary), index=df.index)
    total_ot_salary = pd.Series(np.where(is_office, office_ot_salary, other_ot_salary), index=df.index)

    total_work_on_time_salary = total_workday_salary + lump_sum
    social_security = pd.Series(
        np.where(stated_social_security == 0, 0.0, _social_security_bracket(total_work_on_time_salary)),
        index=df.index,
    )
    total_wage = (
        total_work_on_time_salary + total_sunday_salary + total_ot_salary + diligent_allowance + shift_fee + bonus
    )
    recalculated = total_wage - (social_security + advance_payment + deduct_loan_fund + tax)

    def _masked(series: pd.Series) -> pd.Series:
        return pd.Series(np.where(is_known, series, np.nan), index=df.index)

    # Breakdown columns for the reconciliation report (mirrors the legacy
    # batch's Payroll_Reconcile_Summary layout) — masked to known departments
    # only, since the wage formula (and therefore this breakdown) isn't
    # meaningful for an unrecognized department.
    df["workday_hour"] = _masked(total_workday_hour)
    df["workday_salary"] = _masked(total_workday_salary)
    df["lump_sum"] = _masked(lump_sum)
    df["work_on_time_total"] = _masked(total_work_on_time_salary)
    df["sunday_salary"] = _masked(total_sunday_salary)
    df["ot_salary"] = _masked(total_ot_salary)
    df["diligent_allowance"] = _masked(diligent_allowance)
    df["social_security"] = _masked(social_security)
    df["gross_total"] = _masked(total_wage)
    df["advance_payment"] = _masked(advance_payment)
    df["deduct_loan_fund"] = _masked(deduct_loan_fund)
    df["tax"] = _masked(tax)

    df["stated_amount"] = stated_actual_total_wage
    df["recalculated_amount"] = pd.Series(np.where(is_known, _half_up_round2(recalculated), np.nan), index=df.index)

    stated_rounded = _half_up_round2(stated_actual_total_wage.fillna(0))
    mismatch = is_known & df["recalculated_amount"].notna() & ((stated_rounded - df["recalculated_amount"]).abs() > 1e-6)
    df["variance"] = pd.Series(
        np.where(is_known & df["recalculated_amount"].notna(), stated_rounded - df["recalculated_amount"], np.nan),
        index=df.index,
    )

    data_error = ~df["has_overtime_row"] | df["salary"].isna()
    status = np.select(
        [data_error, ~is_known, mismatch],
        ["dataError", "unrecognizedDepartment", "discrepancy"],
        default="matched",
    )
    df["status"] = status

    detail_map = {
        "dataError": "Missing or unreadable overtime row / attendance data — cannot compute recalculated wage.",
        "unrecognizedDepartment": "Department is not one of the recognized wage-formula categories.",
        "discrepancy": "Recalculated net wage does not match the amount stated in the file.",
        "matched": "",
    }
    df["detail"] = df["status"].map(detail_map)

    return df


# --- Report generation (research.md #6) ---------------------------------


def _num_or_none(value):
    return None if pd.isna(value) else value


def _report_rows(df: pd.DataFrame) -> list[list]:
    rows = []
    for _, row in df.iterrows():
        status = row.get("status")
        rows.append(
            [
                row.get("employee_id"),
                row.get("department"),
                row.get("first_name"),
                row.get("last_name"),
                _num_or_none(row.get("salary")),
                _num_or_none(row.get("total_workday")),
                _num_or_none(row.get("workday_hour")),
                _num_or_none(row.get("total_sunday")),
                _num_or_none(row.get("total_workday_ot_hour")),
                _num_or_none(row.get("total_weekend_ot_hour")),
                _num_or_none(row.get("workday_salary")),
                _num_or_none(row.get("lump_sum")),
                _num_or_none(row.get("work_on_time_total")),
                _num_or_none(row.get("sunday_salary")),
                _num_or_none(row.get("ot_salary")),
                _num_or_none(row.get("diligent_allowance")),
                _num_or_none(row.get("social_security")),
                _num_or_none(row.get("gross_total")),
                _num_or_none(row.get("advance_payment")),
                _num_or_none(row.get("deduct_loan_fund")),
                _num_or_none(row.get("tax")),
                _num_or_none(row.get("stated_amount")),
                _num_or_none(row.get("recalculated_amount")),
                STATUS_LABELS_TH.get(status, status),
                DETAIL_LABELS_TH.get(status, row.get("detail")),
            ]
        )
    return rows


STATUS_COLUMN = REPORT_COLUMNS.index("สถานะ") + 1  # 1-based, for openpyxl cell addressing
UNMATCHED_STATUS_FILL = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")


def _append_rows_with_status_highlight(ws, rows: list[list]) -> None:
    for r in rows:
        ws.append(r)
        if r[STATUS_COLUMN - 1] != STATUS_LABELS_TH["matched"]:
            ws.cell(row=ws.max_row, column=STATUS_COLUMN).fill = UNMATCHED_STATUS_FILL


def build_reconciliation_report(employees: pd.DataFrame) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Reconciliation"
    ws.append(REPORT_COLUMNS)
    _append_rows_with_status_highlight(ws, _report_rows(employees))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_discrepancy_report(discrepancies: pd.DataFrame) -> bytes | None:
    if discrepancies.empty:
        return None
    wb = Workbook()
    ws = wb.active
    ws.title = "Discrepancies"
    ws.append(REPORT_COLUMNS)
    _append_rows_with_status_highlight(ws, _report_rows(discrepancies))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- Orchestrator ---------------------------------------------------------


def _none_if_nan(value):
    return None if value is None or (isinstance(value, float) and pd.isna(value)) else value


def reconcile_payroll(payroll_content: bytes, payroll_filename: str, period: str) -> dict:
    sheet_name = _resolve_sheet_name(period)
    raw = _read_raw_sheet(payroll_content, sheet_name)
    header = _locate_header_positions(raw)
    employees = _build_employee_frame(raw, header)
    employees = _recalculate_wages(employees, header)

    total = len(employees)
    flagged = employees[employees["status"] != "matched"]
    discrepancy_only = employees[employees["status"] == "discrepancy"]

    summary = {
        "total_employees_reconciled": total,
        "matched_count": total - len(flagged),
        "discrepancy_count": len(discrepancy_only),
        "discrepancies_by_department": flagged.groupby("department").size().to_dict() if not flagged.empty else {},
    }

    discrepancies = [
        {
            "department": row["department"],
            "employee_id": row.get("employee_id"),
            "employee_name": (f"{row.get('first_name') or ''} {row.get('last_name') or ''}".strip() or None),
            "stated_amount": _none_if_nan(row.get("stated_amount")),
            "recalculated_amount": _none_if_nan(row.get("recalculated_amount")),
            "variance": _none_if_nan(row.get("variance")),
            "status": row["status"],
            "detail": row["detail"],
        }
        for _, row in flagged.iterrows()
    ]

    reconciliation_bytes = build_reconciliation_report(employees)
    discrepancy_bytes = build_discrepancy_report(discrepancy_only)

    return {
        "summary": summary,
        "discrepancies": discrepancies,
        "reconciliation_report": {
            "filename": "payroll_reconciliation_report.xlsx",
            "content_base64": base64.b64encode(reconciliation_bytes).decode("ascii"),
        },
        "discrepancy_report": (
            {
                "filename": "payroll_discrepancy_report.xlsx",
                "content_base64": base64.b64encode(discrepancy_bytes).decode("ascii"),
            }
            if discrepancy_bytes is not None
            else None
        ),
    }
