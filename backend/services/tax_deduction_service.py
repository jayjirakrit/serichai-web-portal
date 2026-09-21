"""Elderly-Employee Tax Deduction Service (specs/006-elderly-tax-deduction).

Royal Decree No. 639 (B.E. 2560): reads the payroll roster, corrects
Buddhist-Era dates of birth, computes eligibility/ranking/monthly
deductions, and fills the static output template with the results.
"""

import base64
import io
import math
import os
from datetime import date
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from models.common import FileAttachment
from util.be_dates import correct_be_year

AGE_THRESHOLD = 60
SALARY_CAP = 15000.0
HEADCOUNT_CAP_PERCENT = 0.10
HEADCOUNT_CAP_ROUNDING = "floor"
BE_YEAR_OFFSET = 543

DATA_DIR = Path(os.environ.get("DATA_DIR") or (Path(__file__).resolve().parent.parent / "data"))
TEMPLATE_PATH = DATA_DIR / "Tax_Reduction_Template.xlsx"
OUTPUT_REPORT_FILENAME = "tax_deduction_report.xlsx"

ROSTER_SHEET_NAME = "ข้อมูลพนักงาน"
OUTPUT_SHEET_NAME = "ผลประโยช์นพนักงาน"

BONUS_MONTH_OCCURRENCES = {3, 12}

COLUMN_HEADERS: dict[str, str] = {
    "seq": "ลำดับ",
    "id_card_number": "เลขบัตรประชาชน",
    "prefix": "คำนำหน้า",
    "first_name": "ชื่อ",
    "last_name": "สกุล",
    "base_wage_rate": "อัตราค่าแรง",
    "date_of_birth": "วันเดือนปีเกิด",
    "disabled_marker": "คนพิการ",
}

MONTH_NAMES_IN_ORDER = ["ม.ค", "ก.พ", "มี.ค", "เม.ย", "พ.ค", "มิ.ย", "ก.ค", "ส.ค", "ก.ย", "ต.ค", "พ.ย", "ธ.ค"]


class TaxDeductionRequestError(Exception):
    """Raised when the uploaded payroll workbook cannot be processed."""


def _clean_str(value) -> str | None:
    if pd.isna(value):
        return None
    text = str(value).strip()
    return text or None


def _num_or_none(value):
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    if isinstance(value, (np.floating, np.integer)):
        return None if pd.isna(value) else value.item()
    return value


def _int_or_none(value):
    num = _num_or_none(value)
    return None if num is None else int(num)


def read_input(content: bytes) -> pd.DataFrame:
    try:
        raw = pd.read_excel(io.BytesIO(content), sheet_name=ROSTER_SHEET_NAME, header=None, engine="openpyxl")
    except ValueError as exc:
        raise TaxDeductionRequestError(f"Sheet '{ROSTER_SHEET_NAME}' was not found in the uploaded file.") from exc
    except Exception as exc:
        raise TaxDeductionRequestError(
            f"Could not read sheet '{ROSTER_SHEET_NAME}' — file may be empty or unreadable."
        ) from exc

    if raw.empty:
        raise TaxDeductionRequestError(f"Sheet '{ROSTER_SHEET_NAME}' is empty.")

    header = [str(c).strip() if pd.notna(c) else "" for c in raw.iloc[0].tolist()]

    identity_cols: dict[str, int] = {}
    missing_identity: list[str] = []
    for field, label in COLUMN_HEADERS.items():
        if label in header:
            identity_cols[field] = header.index(label)
        else:
            missing_identity.append(label)
    if missing_identity:
        raise TaxDeductionRequestError(
            f"Sheet '{ROSTER_SHEET_NAME}' is missing required column(s): {', '.join(missing_identity)}."
        )

    total_col_positions = [i for i, h in enumerate(header) if h == "รวม"]
    if len(total_col_positions) != 12:
        raise TaxDeductionRequestError(
            f"Sheet '{ROSTER_SHEET_NAME}' must have exactly 12 'รวม' columns (one per month) — "
            f"found {len(total_col_positions)}."
        )

    data = raw.iloc[1:].reset_index(drop=True).copy()
    data.columns = range(len(data.columns))

    df = pd.DataFrame(
        {
            "seq": pd.to_numeric(data[identity_cols["seq"]], errors="coerce"),
            "id_card_number": data[identity_cols["id_card_number"]].apply(_clean_str),
            "prefix": data[identity_cols["prefix"]].apply(_clean_str),
            "first_name": data[identity_cols["first_name"]].apply(_clean_str),
            "last_name": data[identity_cols["last_name"]].apply(_clean_str),
            "base_wage_rate": pd.to_numeric(data[identity_cols["base_wage_rate"]], errors="coerce"),
            "date_of_birth_raw": pd.to_datetime(data[identity_cols["date_of_birth"]], errors="coerce"),
        }
    )

    for month_index, total_col in enumerate(total_col_positions, start=1):
        if month_index in BONUS_MONTH_OCCURRENCES:
            wage_col, ot_col = total_col - 3, total_col - 2
        else:
            wage_col, ot_col = total_col - 2, total_col - 1
        wage = pd.to_numeric(data[wage_col], errors="coerce")
        ot = pd.to_numeric(data[ot_col], errors="coerce")
        df[f"monthly_net_{month_index}"] = wage + ot.fillna(0)

    valid_mask = df["seq"].notna() & df["first_name"].notna()
    df = df[valid_mask].reset_index(drop=True)
    if df.empty:
        raise TaxDeductionRequestError(f"Sheet '{ROSTER_SHEET_NAME}' has no employee rows.")
    df["seq"] = df["seq"].astype(int)

    return df


def calculate(roster: pd.DataFrame, reference_date: date) -> pd.DataFrame:
    df = roster.copy().reset_index(drop=True)

    def _corrected_dob(raw_dob):
        if pd.isna(raw_dob):
            return None
        return correct_be_year(raw_dob.date(), reference_date)

    df["date_of_birth"] = df["date_of_birth_raw"].apply(_corrected_dob)

    def _age(dob):
        if dob is None:
            return np.nan
        years = reference_date.year - dob.year
        if (reference_date.month, reference_date.day) < (dob.month, dob.day):
            years -= 1
        return float(years)

    df["age"] = df["date_of_birth"].apply(_age)

    monthly_cols = [f"monthly_net_{m}" for m in range(1, 13)]
    monthly_values = df[monthly_cols]
    populated_count = monthly_values.notna().sum(axis=1)
    total_populated = monthly_values.sum(axis=1, skipna=True)
    df["average_monthly_net_salary"] = total_populated / populated_count.replace(0, np.nan)

    age_ok = df["age"].notna() & (df["age"] > AGE_THRESHOLD)
    salary_ok = df["average_monthly_net_salary"].notna() & (df["average_monthly_net_salary"] <= SALARY_CAP)
    df["eligible"] = age_ok & salary_ok

    headcount_cap = math.floor(len(df) * HEADCOUNT_CAP_PERCENT)

    df["rank"] = np.nan
    eligible_df = df[df["eligible"]].sort_values(by=["average_monthly_net_salary", "seq"], ascending=[False, True])
    if not eligible_df.empty:
        df.loc[eligible_df.index, "rank"] = range(1, len(eligible_df) + 1)

    df["selected"] = False
    if headcount_cap > 0 and not eligible_df.empty:
        selected_index = df.loc[df["rank"].notna() & (df["rank"] <= headcount_cap)].index
        df.loc[selected_index, "selected"] = True

    for m in range(1, 13):
        net = df[f"monthly_net_{m}"]
        amount = net.where(net <= SALARY_CAP, 0.0).fillna(0.0)
        df[f"monthly_deduction_{m}"] = np.where(df["selected"], amount, 0.0)

    deduction_cols = [f"monthly_deduction_{m}" for m in range(1, 13)]
    df["total_deduction"] = df[deduction_cols].sum(axis=1)
    return df


def fill_output_workbook(results: pd.DataFrame) -> FileAttachment:
    try:
        wb = load_workbook(TEMPLATE_PATH)
    except FileNotFoundError as exc:
        raise TaxDeductionRequestError(f"Output template not found at {TEMPLATE_PATH}.") from exc

    ws = wb.active
    ws.title = OUTPUT_SHEET_NAME

    header_row = [
        str(ws.cell(row=4, column=c).value).strip() if ws.cell(row=4, column=c).value is not None else ""
        for c in range(1, ws.max_column + 1)
    ]

    def col_for(label: str) -> int | None:
        try:
            return header_row.index(label) + 1
        except ValueError:
            return None

    col_seq = col_for("ลำดับ")
    col_id_card = col_for("เลขบัตรประชาชน")
    col_prefix = col_for("คำนำหน้า")
    col_first_name = col_for("ชื่อ")
    col_last_name = col_for("นามสกุล")
    col_salary = col_for("เงินเดือน")
    col_dob = col_for("วันเดือนปีเกิด")
    col_age = col_for("อายุปัจจุบัน")
    month_cols = [col_for(name) for name in MONTH_NAMES_IN_ORDER]
    col_total = col_for("รวม")

    for _, row in results.sort_values("seq").iterrows():
        r = 5 + int(row["seq"]) - 1
        if col_seq:
            ws.cell(row=r, column=col_seq, value=int(row["seq"]))
        if col_id_card:
            ws.cell(row=r, column=col_id_card, value=row.get("id_card_number"))
        if col_prefix:
            ws.cell(row=r, column=col_prefix, value=row.get("prefix"))
        if col_first_name:
            ws.cell(row=r, column=col_first_name, value=row.get("first_name"))
        if col_last_name:
            ws.cell(row=r, column=col_last_name, value=row.get("last_name"))
        if col_salary:
            ws.cell(row=r, column=col_salary, value=_num_or_none(row.get("average_monthly_net_salary")))
        if col_dob:
            ws.cell(row=r, column=col_dob, value=row.get("date_of_birth"))
        if col_age:
            ws.cell(row=r, column=col_age, value=_int_or_none(row.get("age")))
        for month_index, col in enumerate(month_cols, start=1):
            if col:
                ws.cell(row=r, column=col, value=_num_or_none(row.get(f"monthly_deduction_{month_index}")))
        if col_total:
            ws.cell(row=r, column=col_total, value=_num_or_none(row.get("total_deduction")))

    buf = io.BytesIO()
    wb.save(buf)
    return FileAttachment(
        filename=OUTPUT_REPORT_FILENAME, content_base64=base64.b64encode(buf.getvalue()).decode("ascii")
    )


def _employee_result(row: pd.Series) -> dict:
    return {
        "seq": int(row["seq"]),
        "id_card_number": row.get("id_card_number"),
        "prefix": row.get("prefix"),
        "first_name": row.get("first_name"),
        "last_name": row.get("last_name"),
        "date_of_birth": row.get("date_of_birth"),
        "age": _int_or_none(row.get("age")),
        "average_monthly_salary": _num_or_none(row.get("average_monthly_net_salary")),
        "eligible": bool(row["eligible"]),
        "rank": _int_or_none(row.get("rank")),
        "selected": bool(row["selected"]),
        "monthly_amounts": [float(row[f"monthly_deduction_{m}"]) for m in range(1, 13)],
        "total_amount": float(row["total_deduction"]),
    }


def calculate_tax_deduction(payroll_content: bytes, reference_date: date | None) -> dict:
    if reference_date is None:
        reference_date = date.today()

    roster = read_input(payroll_content)
    results = calculate(roster, reference_date)
    report = fill_output_workbook(results)

    ordered = results.sort_values(["total_deduction", "seq"], ascending=[False, True]).reset_index(drop=True)

    summary = {
        "reference_date": reference_date,
        "total_headcount": int(len(ordered)),
        "eligible_count": int(ordered["eligible"].sum()),
        "headcount_cap": math.floor(len(ordered) * HEADCOUNT_CAP_PERCENT),
        "selected_count": int(ordered["selected"].sum()),
        "age_threshold": AGE_THRESHOLD,
        "salary_cap_per_month": SALARY_CAP,
        "headcount_cap_percent": HEADCOUNT_CAP_PERCENT,
        "headcount_cap_rounding": HEADCOUNT_CAP_ROUNDING,
        "be_year_offset": BE_YEAR_OFFSET,
    }

    return {
        "summary": summary,
        "employees": [_employee_result(row) for _, row in ordered.iterrows()],
        "tax_deduction_report": report,
    }
